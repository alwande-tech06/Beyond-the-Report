"""
esgrq.py - ESG Reporting Quality (ESGRQ) index, comparability and assurance engine.

Implements the coding framework in the research paper (sections 9.1-9.5 and Appendix A):
17 disclosure items scored 0-4 per company and report year, grouped into six dimensions.
All inputs come from one Excel workbook (data/esgrq_coding.xlsx) that the group fills in
from the integrated, ESG and assurance reports, with a source document and page for every row.

Nothing in this module invents data: an item without a score is treated as "not yet coded",
never as zero.
"""
from __future__ import annotations

import re
from pathlib import Path

import numpy as np
import pandas as pd

WORKBOOK = Path(__file__).parent / "data" / "esgrq_coding.xlsx"

DIMENSIONS = ["Environmental", "Social", "Governance", "Financial connectivity", "Materiality", "Assurance"]

# code, dimension, indicator, evidence required, <IR> capital
ITEMS = [
    ("E1", "Environmental", "GHG emissions", "Scope 1/2/3 quantities, methodology and trend", "Natural"),
    ("E2", "Environmental", "Energy", "Total energy, renewable energy, targets and progress", "Natural"),
    ("E3", "Environmental", "Water", "Withdrawal/use/recycling, stress context and targets", "Natural"),
    ("E4", "Environmental", "Biodiversity", "Impacts, management and measurable outcomes", "Natural"),
    ("E5", "Environmental", "Tailings", "Facilities, risk, governance and compliance", "Natural"),
    ("E6", "Environmental", "Climate risk", "Physical/transition risks and management", "Natural"),
    ("E7", "Environmental", "Mine closure", "Closure plans, rehabilitation and financial implications", "Natural"),
    ("S1", "Social", "Safety", "Fatalities, TRIFR/LTIFR, targets and trend", "Human"),
    ("S2", "Social", "Workforce", "Employment, diversity, training and wellbeing", "Human"),
    ("S3", "Social", "Community", "Investment, impacts, SLP/CSI and outcomes", "Social and relationship"),
    ("S4", "Social", "Human rights", "Policy, risk assessment, incidents and remedies", "Social and relationship"),
    ("G1", "Governance", "Board ESG oversight", "Committee responsibilities and oversight", "Governance (cross-cutting)"),
    ("G2", "Governance", "Ethics", "Anti-corruption, whistleblowing and incidents", "Governance (cross-cutting)"),
    ("G3", "Governance", "Risk management", "ESG risks integrated into ERM", "Governance (cross-cutting)"),
    ("F1", "Financial connectivity", "Financial connectivity",
     "Cash flow, capex/opex, provisions, impairment, asset value, cost of capital, scenario analysis", "Financial"),
    ("M1", "Materiality", "Materiality", "Process, stakeholder input, financial and impact materiality, prioritisation",
     "Cross-cutting"),
    ("A1", "Assurance", "Assurance", "Provider, standard, scope, level, subject matter and conclusion", "Cross-cutting"),
]
ITEM_TABLE = pd.DataFrame(ITEMS, columns=["code", "dimension", "indicator", "evidence_required", "capital"])

SCALE = {
    0: "Not disclosed",
    1: "Narrative mention only",
    2: "Specific qualitative disclosure or basic quantitative disclosure",
    3: "Quantitative disclosure with target, trend or methodology",
    4: "Quantitative disclosure with target/performance, clear methodology and/or relevant independent assurance",
}

BANDS = [(75, "Strong"), (50, "Moderate"), (-1, "Weak")]

# Common indicators checked for comparability (paper section 9.3)
COMPARABILITY_INDICATORS = [
    "Scope 1 and 2 GHG emissions",
    "Total energy consumption",
    "Water withdrawal",
    "Fatalities",
    "Injury frequency rate (TRIFR or LTIFR)",
    "Socio-economic / community spend",
    "Women in workforce",
]
CRITERIA = {"reported": "Reported by all", "unit": "Same unit", "boundary": "Same boundary",
            "period_end": "Same period end", "rate_basis": "Same rate basis"}

CODING_COLS = ["company", "report_year", "code", "dimension", "indicator", "evidence_required", "score",
               "material", "financial_link", "assured", "evidence_summary", "source_document", "report_page",
               "coder", "score_recheck", "recheck_by", "notes"]
ASSURANCE_COLS = ["company", "report_year", "provider", "standard", "level", "scope", "subject_matter",
                  "conclusion", "source_document", "report_page", "notes"]
COMPARABILITY_COLS = ["indicator", "company", "report_year", "reported", "metric_as_reported", "unit", "boundary",
                      "period_end", "rate_basis", "methodology", "source_document", "report_page", "notes"]

# Claims made in the draft paper. Shown only as prompts to check; never loaded as data.
PAPER_ASSURANCE_CLAIMS = {
    "Exxaro Resources": "Draft paper: KPMG, reasonable assurance over selected KPIs. Verify in the ESG report.",
    "Gold Fields": "Draft paper: reasonable assurance, ISAE 3000 (Revised) and ISAE 3410 for GHG, selected "
                   "information. Verify provider in the assurance report.",
    "Valterra Platinum": "Draft paper: SLR Consulting, AA1000AS v3 Type II, Moderate and High for specified "
                         "information. Verify.",
    "Impala Platinum": "Draft paper: 'AA1000/related' - not specific. Find the assurance statement.",
    "Sibanye-Stillwater": "Draft paper: 'exact coverage to be coded'. Find the assurance statement.",
}


# --------------------------------------------------------------------------------------
# Template
# --------------------------------------------------------------------------------------
def template_frames(companies: list[str], year: int = 2025) -> dict[str, pd.DataFrame]:
    coding = pd.DataFrame(
        [dict(company=c, report_year=year, code=code, dimension=dim, indicator=ind, evidence_required=ev)
         for c in companies for code, dim, ind, ev, _ in ITEMS],
        columns=CODING_COLS)
    assurance = pd.DataFrame(
        [dict(company=c, report_year=year, notes=PAPER_ASSURANCE_CLAIMS.get(c, "")) for c in companies],
        columns=ASSURANCE_COLS)
    comp = pd.DataFrame(
        [dict(indicator=i, company=c, report_year=year) for i in COMPARABILITY_INDICATORS for c in companies],
        columns=COMPARABILITY_COLS)
    return {"Coding": coding, "Assurance": assurance, "Comparability": comp}


def _readme_rows() -> list[list[str]]:
    rows = [["ESGRQ coding workbook - how to fill it in"], [],
            ["1. Coding sheet: one row per company x item. Score each item 0-4 using the scale below."],
            ["2. Every scored row needs source_document and report_page. Unsourced scores are flagged on the dashboard."],
            ["3. material = Yes if the company lists this topic as a material matter. financial_link = Yes if the "
             "report states a financial effect (cost, capex, provision, impairment, revenue, cost of capital)."],
            ["4. assured = Yes if this item's data falls inside the external assurance scope."],
            ["5. score_recheck: a second coder (or the same coder a few days later) re-scores without looking. "
             "The dashboard reports agreement - this is your validation evidence."],
            ["6. Assurance sheet: one row per company, from the independent assurance report. The notes column holds "
             "what the draft paper claims - check it, do not copy it."],
            ["7. Comparability sheet: for each common indicator, record exactly how each company reports it. "
             "Use the same wording for the same thing (e.g. 'operational control'), so the dashboard can match them."],
            ["8. Blank = not yet coded. Never type 0 unless the item is genuinely not disclosed."],
            [], ["Scoring scale"]]
    rows += [[f"{k} = {v}"] for k, v in SCALE.items()]
    rows += [[], ["Items (Appendix A)"]]
    rows += [[f"{code}  {ind} ({dim}): {ev}"] for code, dim, ind, ev, _ in ITEMS]
    return rows


def write_template(path: Path, companies: list[str], year: int = 2025) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    for r in _readme_rows():
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 140

    head_fill = PatternFill("solid", fgColor="DCEBE8")
    input_fill = PatternFill("solid", fgColor="FFF8E1")
    widths = {"company": 20, "indicator": 26, "evidence_required": 44, "evidence_summary": 44, "notes": 50,
              "source_document": 30, "metric_as_reported": 30, "methodology": 30, "subject_matter": 30}
    frames = template_frames(companies, year)
    inputs = {
        "Coding": {"score", "material", "financial_link", "assured", "evidence_summary", "source_document",
                   "report_page", "coder", "score_recheck", "recheck_by", "notes"},
        "Assurance": set(ASSURANCE_COLS) - {"company", "report_year"},
        "Comparability": set(COMPARABILITY_COLS) - {"indicator", "company", "report_year"},
    }
    for name, df in frames.items():
        sh = wb.create_sheet(name)
        sh.append(list(df.columns))
        for row in df.itertuples(index=False):
            sh.append([None if (isinstance(v, float) and np.isnan(v)) else v for v in row])
        for j, col in enumerate(df.columns, start=1):
            cell = sh.cell(row=1, column=j)
            cell.font = Font(bold=True)
            cell.fill = head_fill
            letter = cell.column_letter
            sh.column_dimensions[letter].width = widths.get(col, 14)
            if col in inputs[name]:
                for i in range(2, len(df) + 2):
                    sh.cell(row=i, column=j).fill = input_fill
        sh.freeze_panes = "D2" if name == "Coding" else "C2"
        for r in sh.iter_rows(min_row=2):
            for c in r:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        n = len(df) + 1
        col_letter = {c: sh.cell(row=1, column=j).column_letter for j, c in enumerate(df.columns, start=1)}
        yes_no = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
        sh.add_data_validation(yes_no)
        if name == "Coding":
            dv = DataValidation(type="whole", operator="between", formula1="0", formula2="4", allow_blank=True,
                                showErrorMessage=True, error="Score must be a whole number 0-4")
            sh.add_data_validation(dv)
            for c in ("score", "score_recheck"):
                dv.add(f"{col_letter[c]}2:{col_letter[c]}{n}")
            for c in ("material", "financial_link", "assured"):
                yes_no.add(f"{col_letter[c]}2:{col_letter[c]}{n}")
        elif name == "Comparability":
            yes_no.add(f"{col_letter['reported']}2:{col_letter['reported']}{n}")
    wb.save(path)


# --------------------------------------------------------------------------------------
# Loading
# --------------------------------------------------------------------------------------
def _flag(s: pd.Series) -> pd.Series:
    m = {"yes": 1.0, "y": 1.0, "1": 1.0, "true": 1.0, "no": 0.0, "n": 0.0, "0": 0.0, "false": 0.0}
    return s.map(lambda v: m.get(str(v).strip().lower(), np.nan) if pd.notna(v) else np.nan)


def _text(s: pd.Series) -> pd.Series:
    return s.map(lambda v: "" if pd.isna(v) else str(v).strip())


def load(file) -> dict[str, pd.DataFrame]:
    sheets = pd.read_excel(file, sheet_name=None, dtype=object)
    out = {}
    for name, cols in (("Coding", CODING_COLS), ("Assurance", ASSURANCE_COLS), ("Comparability", COMPARABILITY_COLS)):
        if name not in sheets:
            raise ValueError(f"Workbook has no '{name}' sheet")
        df = sheets[name]
        missing = [c for c in cols if c not in df.columns]
        if missing:
            raise ValueError(f"'{name}' sheet is missing columns: {', '.join(missing)}")
        df = df[cols].dropna(how="all").copy()
        df["report_year"] = pd.to_numeric(df.report_year, errors="coerce").astype("Int64")
        for c in cols:
            if c not in ("report_year", "score", "score_recheck"):
                df[c] = _text(df[c])
        out[name] = df
    c = out["Coding"]
    for col in ("score", "score_recheck"):
        c[col + "_raw"] = c[col]
        c[col] = pd.to_numeric(c[col], errors="coerce")
    for col in ("material", "financial_link", "assured"):
        c[col] = _flag(c[col])
    c["report_page"] = c.report_page.str.replace(r"\.0$", "", regex=True)
    out["Comparability"]["reported"] = _flag(out["Comparability"].reported)
    return out


# --------------------------------------------------------------------------------------
# Scoring
# --------------------------------------------------------------------------------------
def valid_scores(coding: pd.DataFrame) -> pd.DataFrame:
    return coding[coding.score.isin([0, 1, 2, 3, 4])]


def band(v: float) -> str:
    if pd.isna(v):
        return "Not coded"
    return next(label for cut, label in BANDS if v >= cut)


def dimension_scores(coding: pd.DataFrame) -> pd.DataFrame:
    """0-100 per company x year x dimension = mean item score / 4 * 100 over coded items."""
    v = valid_scores(coding)
    if v.empty:
        return pd.DataFrame(columns=["company", "report_year", "dimension", "score", "items_coded", "items_total"])
    g = v.groupby(["company", "report_year", "dimension"])
    d = g.score.mean().mul(25).rename("score").to_frame()
    d["items_coded"] = g.size()
    d = d.reset_index()
    d["items_total"] = d.dimension.map(ITEM_TABLE.groupby("dimension").size())
    return d


def composite(coding: pd.DataFrame, weights: dict[str, float] | None = None) -> pd.DataFrame:
    """ESGRQ composite: weighted mean of dimension scores (equal weights by default)."""
    d = dimension_scores(coding)
    if d.empty:
        return pd.DataFrame(columns=["company", "report_year", *DIMENSIONS, "ESGRQ", "Band", "Coverage"])
    w = pd.Series(weights or {k: 1.0 for k in DIMENSIONS}, dtype=float)
    wide = d.pivot_table(index=["company", "report_year"], columns="dimension", values="score")
    wide = wide.reindex(columns=DIMENSIONS)
    ww = wide.notna().mul(w, axis=1)
    wide["ESGRQ"] = (wide[DIMENSIONS].fillna(0) * w).sum(axis=1) / ww.sum(axis=1).replace(0, np.nan)
    coded = valid_scores(coding).groupby(["company", "report_year"]).size()
    wide["Coverage"] = coded / len(ITEMS) * 100
    wide["Band"] = wide.ESGRQ.map(band)
    return wide.reset_index()


def capital_scores(coding: pd.DataFrame) -> pd.DataFrame:
    v = valid_scores(coding).merge(ITEM_TABLE[["code", "capital"]], on="code", how="left")
    return v.groupby(["company", "report_year", "capital"]).score.mean().mul(25).rename("score").reset_index()


def connectivity(coding: pd.DataFrame) -> pd.DataFrame:
    """Share of coded E/S/G items where the report states a financial effect (supplementary to F1)."""
    v = coding[coding.dimension.isin(["Environmental", "Social", "Governance"]) & coding.financial_link.notna()]
    if v.empty:
        return pd.DataFrame(columns=["company", "report_year", "linked_share", "items"])
    g = v.groupby(["company", "report_year"]).financial_link
    return pd.DataFrame({"linked_share": g.mean() * 100, "items": g.size()}).reset_index()


# --------------------------------------------------------------------------------------
# Comparability (paper 9.3)
# --------------------------------------------------------------------------------------
def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", str(s).strip().lower().replace("₂", "2"))


def comparability(comp: pd.DataFrame, companies: list[str]) -> tuple[pd.DataFrame, float | None]:
    """Per indicator: which criteria hold across companies, and a status.

    Comparable = every criterion holds. Not comparable = fewer than three companies report it, or at most one
    criterion holds. Otherwise partially comparable. Rate basis is only tested where someone filled it in.
    """
    rows = []
    for ind, g in comp[comp.company.isin(companies)].groupby("indicator", sort=False):
        if g.reported.isna().all():
            continue
        rep = g[g.reported == 1]
        res = {"indicator": ind, "companies_reporting": len(rep), "of": len(companies)}
        res["reported"] = len(rep) == len(companies)
        for c in ("unit", "boundary", "period_end", "rate_basis"):
            vals = {_norm(x) for x in rep[c] if str(x).strip()}
            if c == "rate_basis" and not vals:
                res[c] = None
                continue
            filled = rep[c].map(lambda x: bool(str(x).strip())).all()
            res[c] = bool(filled and len(vals) == 1) if len(rep) else False
        tested = [res[c] for c in CRITERIA if res[c] is not None]
        met = sum(tested)
        if met == len(tested):
            status = "Comparable"
        elif len(rep) < 3 or met <= 1:
            status = "Not comparable"
        else:
            status = "Partially comparable"
        res.update(criteria_met=met, criteria_tested=len(tested), status=status)
        rows.append(res)
    out = pd.DataFrame(rows)
    score = (out.status == "Comparable").mean() * 100 if len(out) else None
    return out, score


# --------------------------------------------------------------------------------------
# Validation (paper 13.3)
# --------------------------------------------------------------------------------------
def weighted_kappa(a: pd.Series, b: pd.Series, k: int = 5) -> float:
    """Quadratic-weighted Cohen's kappa for two raters on a 0..k-1 ordinal scale."""
    a, b = a.astype(int).to_numpy(), b.astype(int).to_numpy()
    if len(a) < 2:
        return np.nan
    obs = np.zeros((k, k))
    for x, y in zip(a, b):
        obs[x, y] += 1
    obs /= obs.sum()
    exp = np.outer(obs.sum(1), obs.sum(0))
    i, j = np.indices((k, k))
    w = (i - j) ** 2 / (k - 1) ** 2
    denom = (w * exp).sum()
    return 1 - (w * obs).sum() / denom if denom else np.nan


def validation(coding: pd.DataFrame) -> dict:
    c = coding
    raw = c.score_raw.map(lambda v: "" if pd.isna(v) else str(v).strip())
    invalid = c[(raw != "") & ~c.score.isin([0, 1, 2, 3, 4])]
    scored = valid_scores(c)
    unsourced = scored[(scored.source_document == "") | (scored.report_page == "")]
    dupes = c[c.duplicated(["company", "report_year", "code"], keep=False)]
    unknown = c[~c.code.isin(ITEM_TABLE.code)]
    pair = scored[scored.score_recheck.isin([0, 1, 2, 3, 4])]
    disagree = pair[pair.score != pair.score_recheck]
    return dict(
        expected=len(c), scored=len(scored),
        traceable=len(scored) - len(unsourced), unsourced=unsourced,
        invalid=invalid, duplicates=dupes, unknown=unknown,
        rechecked=len(pair),
        exact=(pair.score == pair.score_recheck).mean() * 100 if len(pair) else np.nan,
        within1=((pair.score - pair.score_recheck).abs() <= 1).mean() * 100 if len(pair) else np.nan,
        kappa=weighted_kappa(pair.score, pair.score_recheck) if len(pair) >= 2 else np.nan,
        disagreements=disagree,
    )


# --------------------------------------------------------------------------------------
# Blind second-coder recheck (paper 13.3)
# --------------------------------------------------------------------------------------
RECHECK = WORKBOOK.with_name("esgrq_recheck_blind.xlsx")
BLIND_COLS = ["company", "report_year", "code", "dimension", "indicator", "evidence_required", "source_document",
              "report_page", "score_recheck", "recheck_by", "recheck_notes"]


def write_blind(coding: pd.DataFrame, path: Path = RECHECK) -> int:
    """Recheck sheet without the first coder's score or evidence summary, so the second coder scores independently."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    rows = valid_scores(coding).copy()
    rows["recheck_notes"] = ""
    rows["recheck_by"] = ""
    rows["score_recheck"] = np.nan
    blind = rows[BLIND_COLS]
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame({"How to recheck": [
            "Do NOT open esgrq_coding.xlsx until you have finished.",
            "For each row, open the source document at the page(s) given, read the evidence yourself,",
            "and enter a whole-number score 0-4 in score_recheck using the scale below. Put your name in recheck_by.",
            "Then run:  python esgrq.py merge   (copies your scores into esgrq_coding.xlsx for the Validation tab).",
            "", *[f"{k} = {v}" for k, v in SCALE.items()]]}).to_excel(xw, sheet_name="Read me", index=False)
        blind.to_excel(xw, sheet_name="Recheck", index=False)
    wb = load_workbook(path)
    ws = wb["Recheck"]
    col = {c.value: c.column_letter for c in ws[1]}
    for c in ws[1]:
        c.font = Font(bold=True)
    fill = PatternFill("solid", fgColor="FFF8E1")
    for r in range(2, len(blind) + 2):
        for k in ("score_recheck", "recheck_by", "recheck_notes"):
            ws[f"{col[k]}{r}"].fill = fill
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="4", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col['score_recheck']}2:{col['score_recheck']}{len(blind) + 1}")
    for k, w in {"indicator": 24, "evidence_required": 50, "source_document": 40, "recheck_notes": 40}.items():
        ws.column_dimensions[col[k]].width = w
    ws.freeze_panes = "D2"
    wb["Read me"].column_dimensions["A"].width = 110
    wb.save(path)
    return len(blind)


def merge_recheck(path: Path = RECHECK, target: Path = WORKBOOK) -> int:
    from openpyxl import load_workbook

    r = pd.read_excel(path, sheet_name="Recheck")
    r = r[pd.to_numeric(r.score_recheck, errors="coerce").isin([0, 1, 2, 3, 4])]
    key = {(a.company, int(a.report_year), a.code): a for a in r.itertuples()}
    wb = load_workbook(target)
    ws = wb["Coding"]
    h = {c.value: c.column for c in ws[1]}
    n = 0
    for i in range(2, ws.max_row + 1):
        k = (ws.cell(i, h["company"]).value, ws.cell(i, h["report_year"]).value, ws.cell(i, h["code"]).value)
        if k in key:
            ws.cell(i, h["score_recheck"]).value = int(key[k].score_recheck)
            ws.cell(i, h["recheck_by"]).value = None if pd.isna(key[k].recheck_by) else key[k].recheck_by
            n += 1
    wb.save(target)
    return n


if __name__ == "__main__":
    import sys

    import loader as L

    if sys.argv[1:] == ["blind"]:
        print(f"Wrote {write_blind(load(WORKBOOK)['Coding'])} rows to {RECHECK}")
        sys.exit()
    if sys.argv[1:] == ["merge"]:
        print(f"Merged {merge_recheck()} rechecked scores into {WORKBOOK}")
        sys.exit()
    if not WORKBOOK.exists():
        write_template(WORKBOOK, list(L.COMPANIES))
        print(f"Wrote blank template to {WORKBOOK}")
    data = load(WORKBOOK)
    pd.set_option("display.width", 220)
    print(composite(data["Coding"]).round(1).to_string())
    print(comparability(data["Comparability"], list(L.COMPANIES))[0].to_string())
    v = validation(data["Coding"])
    print({k: v[k] for k in ("expected", "scored", "traceable", "rechecked", "exact", "within1", "kappa")})
