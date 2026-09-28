"""
loader.py - extraction, standardisation and ratio engine for the APFA802 dashboard.

Every financial number is read directly from the group's Excel workbooks (no hand
re-typing), and every observation keeps a pointer back to the file, sheet, line item
and Excel row it came from. Page numbers in the integrated report still have to be
added by the group (column `report_page`) to meet the brief's traceability rule.
"""
from __future__ import annotations

import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

DATA_DIR = Path(__file__).parent / "data"
YEARS = [2021, 2022, 2023, 2024, 2025]

# --------------------------------------------------------------------------------------
# Company register
# --------------------------------------------------------------------------------------
COMPANIES = {
    "Valterra Platinum": dict(
        file="valterra_platinum.xlsx", ticker="VAL", commodity="PGMs",
        currency="ZAR", fy_end="31 Dec", scale="suffix",  # values like 215B / 966M in rand
        colour="#1F3B5A",
    ),
    "Impala Platinum": dict(
        file="impala_platinum.xlsx", ticker="IMP", commodity="PGMs",
        currency="ZAR", fy_end="30 Jun", scale="Rm", colour="#F39C12",
    ),
    "Sibanye-Stillwater": dict(
        file="sibanye_stillwater.xlsx", ticker="SSW", commodity="PGMs & gold",
        currency="ZAR", fy_end="31 Dec", scale="suffix", colour="#5DA9E9",
    ),
    "Gold Fields": dict(
        file="gold_fields.xlsx", ticker="GFI", commodity="Gold",
        currency="USD", fy_end="31 Dec", scale="USDm", colour="#8C98A4",
    ),
    "Exxaro Resources": dict(
        file="exxaro_resources.xlsx", ticker="EXX", commodity="Coal & iron ore",
        currency="ZAR", fy_end="31 Dec", scale="Rm", colour="#2E8B80",
    ),
}

# --------------------------------------------------------------------------------------
# Line-item mapping: metric -> list of (sheet, label, occurrence, sign)
# occurrence = which instance of a repeated label (0 = first); sign converts the
# company's presentation into the dashboard convention:
#   costs, capex, dividends, finance costs and tax expense are stored as POSITIVE numbers.
# --------------------------------------------------------------------------------------
IS, FP, CF = "IS", "FP", "CF"

SHEETS = {
    "Valterra Platinum": {IS: "Income statement", FP: "Financial position", CF: "Cash Flow"},
    "Impala Platinum": {IS: "Income Statement", FP: "Financial Position", CF: "Cash Flow Statement"},
    "Sibanye-Stillwater": {IS: "Income Statements", FP: "Blance Sheet", CF: "Cash Flow"},
    "Gold Fields": {IS: "Income Statement", FP: "SOFP", CF: "SOCF"},
    "Exxaro Resources": {IS: "Income Statement", FP: "Financial Position", CF: "Cash flow"},
}

MAP: dict[str, dict[str, list[tuple]]] = {
    "Valterra Platinum": {
        "revenue": [(IS, "Total Revenues", 0, 1)],
        "cost_of_sales": [(IS, "Cost of Goods Sold, Total", 0, 1)],
        "finance_costs": [(IS, "Interest Expense, Total", 0, -1)],
        "profit_before_tax": [(IS, "EBT, Incl. Unusual Items", 0, 1)],
        "tax_expense": [(IS, "Income Tax Expense", 0, 1)],
        "net_profit": [(IS, "Net Income to Company", 0, 1)],
        "total_assets": [(FP, "Total Assets", 0, 1)],
        "total_equity": [(FP, "Total Equity", 0, 1)],
        "total_liabilities": [(FP, "Total Liabilities", 0, 1)],
        "current_assets": [(FP, "Total Current Assets", 0, 1)],
        "current_liabilities": [(FP, "Total Current Liabilities", 0, 1)],
        "inventory": [(FP, "Inventory", 0, 1)],
        "cash": [(FP, "Cash And Equivalents", 0, 1)],
        "ppe": [(FP, "Net Property Plant And Equipment", 0, 1)],
        "total_debt": [(FP, "Total Debt", 0, 1)],
        "operating_cash_flow": [(CF, "Cash from Operations", 0, 1)],
        "capex": [(CF, "Capital Expenditure", 0, -1)],
        "dividends_paid": [(CF, "Common Dividends Paid", 0, -1), (CF, "Special Dividend Paid", 0, -1)],
    },
    "Sibanye-Stillwater": {
        "revenue": [(IS, "Total Revenues", 0, 1)],
        "cost_of_sales": [(IS, "Cost of Goods Sold, Total", 0, 1)],
        "finance_costs": [(IS, "Interest Expense, Total", 0, -1)],
        "profit_before_tax": [(IS, "EBT, Incl. Unusual Items", 0, 1)],
        "tax_expense": [(IS, "Income Tax Expense", 0, 1)],
        "net_profit": [(IS, "Net Income to Company", 0, 1)],
        "total_assets": [(FP, "Total Assets", 0, 1)],
        "total_equity": [(FP, "Total Equity", 0, 1)],
        "total_liabilities": [(FP, "Total Liabilities", 0, 1)],
        "current_assets": [(FP, "Total Current Assets", 0, 1)],
        "current_liabilities": [(FP, "Total Current Liabilities", 0, 1)],
        "inventory": [(FP, "Inventory", 0, 1)],
        "cash": [(FP, "Cash And Equivalents", 0, 1)],
        "ppe": [(FP, "Net Property Plant And Equipment", 0, 1)],
        "total_debt": [(FP, "Total Debt", 0, 1)],
        "operating_cash_flow": [(CF, "Cash from Operations", 0, 1)],
        "capex": [(CF, "Capital Expenditure", 0, -1)],
        "dividends_paid": [(CF, "Common Dividends Paid", 0, -1)],
    },
    "Impala Platinum": {
        "revenue": [(IS, "Revenue", 0, 1)],
        "cost_of_sales": [(IS, "Cost of sales", 0, -1)],
        "finance_costs": [(IS, "Finance costs", 0, -1)],
        "profit_before_tax": [(IS, "Profit/(loss) before tax", 0, 1)],
        "tax_expense": [(IS, "Income tax (expense)/credit", 0, -1)],
        "net_profit": [(IS, "Profit/(loss) for the year", 0, 1)],
        "total_assets": [(FP, "Total assets", 0, 1)],
        "total_equity": [(FP, "Total equity", 0, 1)],
        "total_liabilities": [(FP, "Total liabilities", 0, 1)],
        "current_assets": [(FP, "Current assets", 0, 1)],
        "current_liabilities": [(FP, "Current liabilities", 0, 1)],
        "inventory": [(FP, "Inventories", 0, 1)],
        "cash": [(FP, "Cash and cash equivalents", 0, 1)],
        "ppe": [(FP, "Property, plant and equipment", 0, 1)],
        "total_debt": [(FP, "Borrowings", 0, 1), (FP, "Borrowings", 1, 1)],
        "operating_cash_flow": [(CF, "Net cash inflow from operating activities", 0, 1)],
        "capex": [(CF, "Purchase of property, plant and equipment", 0, -1)],
        "dividends_paid": [(CF, "Dividends paid to shareholders of the Company", 0, -1)],
    },
    "Gold Fields": {
        "revenue": [(IS, "Revenue", 0, 1)],
        "cost_of_sales": [(IS, "Cost of sales", 0, -1)],
        "finance_costs": [(IS, "Net interest expense", 0, -1)],
        "profit_before_tax": [(IS, "Profit before taxation", 0, 1)],
        "tax_expense": [(IS, "Mining and income taxation", 0, -1)],
        "net_profit": [(IS, "Profit for the year", 0, 1)],
        "total_assets": [(FP, "Total assets", 0, 1)],
        "total_equity": [(FP, "Total equity", 0, 1)],
        "total_liabilities": [(FP, "Non-current liabilities", 0, 1), (FP, "Current liabilities", 0, 1)],
        "current_assets": [(FP, "Current assets", 0, 1)],
        "current_liabilities": [(FP, "Current liabilities", 0, 1)],
        "cash": [(FP, "Cash and cash equivalents", 0, 1)],
        "ppe": [(FP, "Property, plant and equipment", 0, 1)],
        "total_debt": [
            (FP, "Borrowings", 0, 1), (FP, "Current portion of borrowings", 0, 1),
            (FP, "Lease liabilities", 0, 1), (FP, "Current portion of lease liabilities", 0, 1),
        ],
        "operating_cash_flow": [(CF, "Cash flows from operating activities", 0, 1)],
        "capex": [(CF, "Capital expenditure – additions", 0, -1)],
        "dividends_paid": [(CF, "Owners of the parent", 0, -1)],
    },
    "Exxaro Resources": {
        "revenue": [(IS, "Revenue (note 7)", 0, 1)],
        "cost_of_sales": [],  # Exxaro presents operating expenses by nature, no cost of sales line
        "operating_expenses": [(IS, "Operating expenses (note 8)", 0, -1)],
        "finance_costs": [(IS, "Finance costs (note 10)", 0, -1)],
        "profit_before_tax": [(IS, "Profit before tax", 0, 1)],
        "tax_expense": [(IS, "Income tax expense", 0, -1)],
        "net_profit": [(IS, "Profit for the year", 0, 1)],
        "total_assets": [(FP, "Total assets", 0, 1)],
        "total_equity": [(FP, "Total equity", 0, 1)],
        "total_liabilities": [(FP, "Total liabilities", 0, 1)],
        "current_assets": [(FP, "Current assets", 0, 1)],
        "current_liabilities": [(FP, "Current liabilities", 0, 1)],
        "inventory": [(FP, "Inventories", 0, 1)],
        "cash": [(FP, "Cash and cash equivalents (note 20)", 0, 1)],
        "ppe": [(FP, "Property, plant and equipment", 0, 1)],
        "total_debt": [
            (FP, "Interest-bearing borrowings (note 15)", 0, 1), (FP, "Interest-bearing borrowings (note 15)", 1, 1),
            (FP, "Lease liabilities (note 16)", 0, 1), (FP, "Lease liabilities (note 16)", 1, 1),
            (FP, "Overdraft (note 17)", 0, 1),
        ],
        "operating_cash_flow": [(CF, "Cash flows from operating activities", 0, 1)],
        "capex": [(CF, "Property, plant and equipment acquired (note 12)", 0, -1)],
        "dividends_paid": [(CF, "Dividends paid to owners of the parent (note 5)", 0, -1)],
    },
}

# --------------------------------------------------------------------------------------
# Data dictionary
# --------------------------------------------------------------------------------------
DICTIONARY = [
    # name, label, type, unit, definition / formula, better
    ("revenue", "Revenue", "Extracted", "currency m", "Total revenue as reported in the income statement.", "higher"),
    ("cost_of_sales", "Cost of sales", "Extracted", "currency m", "Cost of sales / cost of goods sold (positive). Not reported by Exxaro.", "lower"),
    ("operating_expenses", "Operating expenses", "Extracted", "currency m", "Exxaro only: operating expenses by nature (positive).", "lower"),
    ("finance_costs", "Finance costs", "Extracted", "currency m", "Interest expense (Gold Fields: net interest expense), positive.", "lower"),
    ("profit_before_tax", "Profit before tax", "Extracted", "currency m", "Profit/(loss) before tax including unusual items.", "higher"),
    ("tax_expense", "Tax expense", "Extracted", "currency m", "Income/mining tax expense; negative = tax credit.", "n/a"),
    ("net_profit", "Net profit", "Extracted", "currency m", "Profit/(loss) for the year before non-controlling interests.", "higher"),
    ("total_assets", "Total assets", "Extracted", "currency m", "Total assets at year end.", "n/a"),
    ("total_equity", "Total equity", "Extracted", "currency m", "Total equity incl. non-controlling interests.", "higher"),
    ("total_liabilities", "Total liabilities", "Extracted", "currency m", "Total liabilities (Gold Fields: non-current + current).", "lower"),
    ("current_assets", "Current assets", "Extracted", "currency m", "Total current assets.", "n/a"),
    ("current_liabilities", "Current liabilities", "Extracted", "currency m", "Total current liabilities.", "n/a"),
    ("inventory", "Inventory", "Extracted", "currency m", "Inventories. Not separately captured for Gold Fields.", "n/a"),
    ("cash", "Cash", "Extracted", "currency m", "Cash and cash equivalents.", "higher"),
    ("ppe", "Property, plant & equipment", "Extracted", "currency m", "Net PPE (manufactured capital).", "n/a"),
    ("total_debt", "Total debt", "Extracted", "currency m", "Interest-bearing borrowings + lease liabilities (current + non-current).", "lower"),
    ("operating_cash_flow", "Operating cash flow", "Extracted", "currency m", "Net cash from operating activities.", "higher"),
    ("capex", "Capital expenditure", "Extracted", "currency m", "Purchase of PPE / capital additions (positive).", "n/a"),
    ("dividends_paid", "Dividends paid", "Extracted", "currency m", "Dividends paid to owners of the parent incl. special dividends (positive).", "n/a"),
    ("gross_profit", "Gross profit", "Derived", "currency m", "revenue - cost_of_sales", "higher"),
    ("ebit", "EBIT (proxy)", "Derived", "currency m", "profit_before_tax + finance_costs. Used instead of reported operating profit so all five firms are measured the same way.", "higher"),
    ("free_cash_flow", "Free cash flow", "Derived", "currency m", "operating_cash_flow - capex", "higher"),
    ("net_debt", "Net debt", "Derived", "currency m", "total_debt - cash (negative = net cash)", "lower"),
    ("revenue_growth", "Revenue growth", "Ratio", "%", "revenue_t / revenue_t-1 - 1", "higher"),
    ("gross_margin", "Gross margin", "Ratio", "%", "gross_profit / revenue", "higher"),
    ("ebit_margin", "EBIT margin", "Ratio", "%", "ebit / revenue", "higher"),
    ("net_margin", "Net margin", "Ratio", "%", "net_profit / revenue", "higher"),
    ("roa", "Return on assets", "Ratio", "%", "net_profit / total_assets (year end)", "higher"),
    ("roe", "Return on equity", "Ratio", "%", "net_profit / total_equity (year end)", "higher"),
    ("current_ratio", "Current ratio", "Ratio", "x", "current_assets / current_liabilities", "higher"),
    ("quick_ratio", "Quick ratio", "Ratio", "x", "(current_assets - inventory) / current_liabilities", "higher"),
    ("debt_to_equity", "Debt to equity", "Ratio", "x", "total_debt / total_equity", "lower"),
    ("liabilities_to_assets", "Liabilities to assets", "Ratio", "%", "total_liabilities / total_assets", "lower"),
    ("interest_cover", "Interest cover", "Ratio", "x", "ebit / finance_costs", "higher"),
    ("ocf_margin", "Operating cash flow margin", "Ratio", "%", "operating_cash_flow / revenue", "higher"),
    ("fcf_margin", "Free cash flow margin", "Ratio", "%", "free_cash_flow / revenue", "higher"),
    ("capex_intensity", "Capex intensity", "Ratio", "%", "capex / revenue (reinvestment in manufactured capital)", "n/a"),
    ("effective_tax_rate", "Effective tax rate", "Ratio", "%", "tax_expense / profit_before_tax (blank when PBT <= 0)", "n/a"),
    ("dividends_to_ocf", "Dividends / operating cash flow", "Ratio", "%", "dividends_paid / operating_cash_flow", "n/a"),
]
DICT_DF = pd.DataFrame(DICTIONARY, columns=["variable", "label", "type", "unit", "definition", "better_when"])
LABELS = dict(zip(DICT_DF.variable, DICT_DF.label))
UNITS = dict(zip(DICT_DF.variable, DICT_DF.unit))
RATIOS = DICT_DF.loc[DICT_DF.type == "Ratio", "variable"].tolist()
AMOUNTS = DICT_DF.loc[DICT_DF.type != "Ratio", "variable"].tolist()

# --------------------------------------------------------------------------------------
# Parsing helpers
# --------------------------------------------------------------------------------------
_SUFFIX = {"K": 1e3, "M": 1e6, "B": 1e9, "T": 1e12}


def norm(label: str) -> str:
    return re.sub(r"\s+", " ", str(label)).strip().lower()


def parse_value(v, scale: str) -> float:
    """Convert a cell value to a float in millions."""
    if v is None:
        return np.nan
    if isinstance(v, (int, float)):
        return float(v) / (1e6 if scale == "suffix" else 1)
    s = str(v).strip().replace("\u00a0", " ")
    if s in {"", "-", "—", "–", "n/a", "NA"}:
        return np.nan
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()").replace(",", "").replace(" ", "")
    mult = 1.0
    if s and s[-1].upper() in _SUFFIX:
        mult = _SUFFIX[s[-1].upper()]
        s = s[:-1]
    try:
        x = float(s) * mult
    except ValueError:
        return np.nan
    x = -x if neg else x
    if scale == "suffix":
        x /= 1e6  # rand -> R million
    return x


def read_sheet(path: Path, sheet: str, scale: str) -> dict:
    """Return {(normalised label, occurrence): {"label", "row", year: value}}."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    year_cols, header_idx = {}, None
    for i, row in enumerate(rows[:10]):
        found = {}
        for j, c in enumerate(row):
            if c is None:
                continue
            m = re.search(r"(20\d\d)", str(c))
            if m and len(str(c).strip()) <= 30:
                found[j] = int(m.group(1))
        if len(found) >= 3:
            year_cols, header_idx = found, i
            break
    if header_idx is None:
        raise ValueError(f"No year header found in {path.name}:{sheet}")

    first_year_col = min(year_cols)
    out, seen = {}, {}
    for i, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        label = None
        for j in range(min(first_year_col, len(row))):
            if isinstance(row[j], str) and row[j].strip():
                label = row[j]
                break
        if label is None:
            continue
        key = norm(label)
        occ = seen.get(key, 0)
        seen[key] = occ + 1
        rec = {"label": re.sub(r"\s+", " ", label).strip(), "row": i}
        for j, yr in year_cols.items():
            rec[yr] = parse_value(row[j] if j < len(row) else None, scale)
        out[(key, occ)] = rec
    return out


# --------------------------------------------------------------------------------------
# Build dataset
# --------------------------------------------------------------------------------------
def build_long() -> pd.DataFrame:
    """Tidy, traceable dataset of extracted line items (one row per company-year-metric)."""
    records = []
    for company, meta in COMPANIES.items():
        path = DATA_DIR / meta["file"]
        cache = {k: read_sheet(path, s, meta["scale"]) for k, s in SHEETS[company].items()}
        for metric, parts in MAP[company].items():
            if not parts:
                continue
            for yr in YEARS:
                vals, sources = [], []
                for sheet_key, label, occ, sign in parts:
                    rec = cache[sheet_key].get((norm(label), occ))
                    if rec is None:
                        raise KeyError(f"{company}: '{label}' (occurrence {occ}) not found in {sheet_key}")
                    v = rec.get(yr, np.nan)
                    if not np.isnan(v):
                        vals.append(v * sign)
                    sources.append(f"{SHEETS[company][sheet_key]}!row {rec['row']} '{rec['label']}'")
                value = float(np.sum(vals)) if vals else np.nan
                records.append(dict(
                    company=company, ticker=meta["ticker"], fiscal_year=yr, fy_end=meta["fy_end"],
                    variable=metric, value_m=value, currency=meta["currency"],
                    source_file=meta["file"], source_lines=" + ".join(sources), report_page="",
                ))
    return pd.DataFrame(records)


def _safe_div(a, b):
    with np.errstate(divide="ignore", invalid="ignore"):
        r = a / b
    return r.replace([np.inf, -np.inf], np.nan)


def build_wide(long: pd.DataFrame, fx: dict | None = None) -> pd.DataFrame:
    """Company-year panel with derived amounts and ratios.

    fx: optional {year: ZAR per USD}; when given, USD amounts are converted to R million
    (ratios are unit-free and unaffected).
    """
    w = long.pivot_table(index=["company", "fiscal_year"], columns="variable", values="value_m", aggfunc="first")
    w = w.reset_index()
    for col in AMOUNTS:
        if col not in w and DICT_DF.set_index("variable").loc[col, "type"] == "Extracted":
            w[col] = np.nan
    w["currency"] = w.company.map(lambda c: COMPANIES[c]["currency"])

    if fx:
        usd = w.currency == "USD"
        rate = w.fiscal_year.map(fx)
        for col in [c for c in w.columns if c in AMOUNTS]:
            w.loc[usd, col] = w.loc[usd, col] * rate[usd]
        w.loc[usd, "currency"] = "ZAR (converted)"

    w["gross_profit"] = w.revenue - w.cost_of_sales
    w["ebit"] = w.profit_before_tax + w.finance_costs
    w["free_cash_flow"] = w.operating_cash_flow - w.capex
    w["net_debt"] = w.total_debt - w.cash

    w = w.sort_values(["company", "fiscal_year"])
    w["revenue_growth"] = w.groupby("company").revenue.pct_change() * 100
    w["gross_margin"] = _safe_div(w.gross_profit, w.revenue) * 100
    w["ebit_margin"] = _safe_div(w.ebit, w.revenue) * 100
    w["net_margin"] = _safe_div(w.net_profit, w.revenue) * 100
    w["roa"] = _safe_div(w.net_profit, w.total_assets) * 100
    w["roe"] = _safe_div(w.net_profit, w.total_equity) * 100
    w["current_ratio"] = _safe_div(w.current_assets, w.current_liabilities)
    w["quick_ratio"] = _safe_div(w.current_assets - w.inventory.fillna(0), w.current_liabilities)
    w.loc[w.inventory.isna(), "quick_ratio"] = np.nan
    w["debt_to_equity"] = _safe_div(w.total_debt, w.total_equity)
    w["liabilities_to_assets"] = _safe_div(w.total_liabilities, w.total_assets) * 100
    w["interest_cover"] = _safe_div(w.ebit, w.finance_costs)
    w["ocf_margin"] = _safe_div(w.operating_cash_flow, w.revenue) * 100
    w["fcf_margin"] = _safe_div(w.free_cash_flow, w.revenue) * 100
    w["capex_intensity"] = _safe_div(w.capex, w.revenue) * 100
    w["effective_tax_rate"] = np.where(w.profit_before_tax > 0, _safe_div(w.tax_expense, w.profit_before_tax) * 100, np.nan)
    w["dividends_to_ocf"] = _safe_div(w.dividends_paid, w.operating_cash_flow) * 100
    return w.reset_index(drop=True)


# --------------------------------------------------------------------------------------
# Data-quality checks (anomaly detection on the extraction itself)
# --------------------------------------------------------------------------------------
def quality_checks(wide: pd.DataFrame, long: pd.DataFrame) -> pd.DataFrame:
    issues = []
    for _, r in wide.iterrows():
        c, y = r.company, r.fiscal_year
        # 1. accounting identity
        if pd.notna(r.total_assets) and pd.notna(r.total_equity) and pd.notna(r.total_liabilities):
            gap = r.total_assets - (r.total_equity + r.total_liabilities)
            if abs(gap) > 0.01 * abs(r.total_assets):
                issues.append((c, y, "Balance sheet does not balance",
                               f"Assets - (equity + liabilities) = {gap:,.0f}m ({gap / r.total_assets:.1%})", "High"))
        # 2. profit sign vs cash
        if pd.notna(r.net_profit) and pd.notna(r.operating_cash_flow) and r.net_profit < 0 < r.operating_cash_flow:
            issues.append((c, y, "Loss year with positive operating cash flow",
                           "Usually non-cash impairments; check the impairment note for the driver.", "Info"))
        # 3. tax credit / unusual rate
        if pd.notna(r.effective_tax_rate) and not (0 <= r.effective_tax_rate <= 60):
            issues.append((c, y, "Unusual effective tax rate", f"{r.effective_tax_rate:.1f}% (SA statutory rate is 27%)", "Medium"))

    # 4. large year-on-year jumps in extracted line items (possible capture errors)
    for (c, v), g in long.groupby(["company", "variable"]):
        g = g.sort_values("fiscal_year")
        vals = g.value_m.values
        for i in range(1, len(vals)):
            a, b = vals[i - 1], vals[i]
            if np.isnan(a) or np.isnan(b) or abs(a) < 1:
                continue
            ch = (b - a) / abs(a)
            if abs(ch) > 1.5 and v in {"revenue", "total_assets", "total_equity", "cash", "ppe", "current_assets"}:
                issues.append((c, int(g.fiscal_year.values[i]), f"Large change in {LABELS.get(v, v).lower()}",
                               f"{ch:+.0%} year on year; confirm against the report.", "Medium"))

    # 5. missing values
    miss = long[long.value_m.isna()]
    for (c, v), g in miss.groupby(["company", "variable"]):
        yrs = ", ".join(str(y) for y in sorted(g.fiscal_year))
        issues.append((c, None, f"Missing: {LABELS.get(v, v).lower()}", f"No value for {yrs}", "Low"))

    # 6. known issues spotted while building the dataset
    known = [
        ("Gold Fields", None, "Reported in US dollars",
         "Gold Fields reports in US$ million. Ratios are comparable; amounts need the ZAR conversion toggle.", "Info"),
        ("Impala Platinum", None, "June year end",
         "Impala's financial year ends 30 June, the others 31 December. FY2025 = July 2024 to June 2025.", "Info"),
        ("Gold Fields", 2025, "Capture error in workbook",
         "'Other non-current assets' 2025 (11,336.8) repeats the PPE figure; several 2025 cash-flow and cost lines are 0 instead of blank.", "High"),
        ("Exxaro Resources", 2022, "Mislabelled rows in workbook",
         "2021-2022 figures sit on the wrong labels for equity-accounted income ('impairment' row holds 9,790 / 6,204). Profit for the year is correct.", "Medium"),
        ("Exxaro Resources", None, "No cost of sales",
         "Exxaro presents expenses by nature, so gross margin is not available. Use EBIT margin for comparison.", "Info"),
        ("Valterra Platinum", None, "Data-provider figures, rounded",
         "Workbook uses a data-provider layout (e.g. '215B'), rounded to 3 significant figures. Re-capture key lines from the integrated report with page numbers.", "Medium"),
        ("Sibanye-Stillwater", None, "Data-provider figures, rounded",
         "Workbook uses a data-provider layout (e.g. '172B'), rounded to 3 significant figures. Re-capture key lines from the integrated report with page numbers.", "Medium"),
        ("Exxaro Resources", 2023, "Cash does not roll forward",
         "Closing cash 2022 is 7,041 but opening cash 2023 is 14,812; 2023 'dividends received' (4,911) is repeated on two lines. Check against the 2023 report.", "High"),
        ("Sibanye-Stillwater", 2025, "Incomplete 2025 extraction",
         "Several 2025 cash-flow detail lines and PPE breakdowns are '-'. Totals used by the dashboard are present.", "Low"),
    ]
    issues.extend(known)
    df = pd.DataFrame(issues, columns=["company", "fiscal_year", "check", "detail", "severity"])
    order = {"High": 0, "Medium": 1, "Low": 2, "Info": 3}
    return df.sort_values(["severity", "company"], key=lambda s: s.map(order) if s.name == "severity" else s).reset_index(drop=True)


# --------------------------------------------------------------------------------------
# Risk scorecard
# --------------------------------------------------------------------------------------
PILLARS = {
    "Profitability": [("net_margin", 1), ("roa", 1)],
    "Liquidity": [("current_ratio", 1), ("quick_ratio", 1)],
    "Solvency": [("debt_to_equity", -1), ("liabilities_to_assets", -1), ("interest_cover", 1)],
    "Cash generation": [("ocf_margin", 1), ("fcf_margin", 1)],
}
CAPS = {"interest_cover": (-20, 50), "current_ratio": (0, 5), "quick_ratio": (0, 4), "debt_to_equity": (0, 1.5),
        "net_margin": (-40, 50), "roa": (-30, 40), "ocf_margin": (-10, 60), "fcf_margin": (-30, 50)}


def minmax_score(s: pd.Series, direction: int, cap=None) -> pd.Series:
    x = s.clip(*cap) if cap else s
    lo, hi = x.min(), x.max()
    if pd.isna(lo) or hi == lo:
        return pd.Series(50.0, index=s.index).where(s.notna())
    sc = (x - lo) / (hi - lo) * 100
    return sc if direction > 0 else 100 - sc


def scorecard(wide: pd.DataFrame, esg_scores: pd.DataFrame | None = None, weights: dict | None = None) -> pd.DataFrame:
    """Pillar scores 0-100 (100 = strongest / lowest risk), scaled across all company-years
    so a firm's score is comparable across time as well as across peers."""
    sc = wide[["company", "fiscal_year"]].copy()
    for pillar, items in PILLARS.items():
        parts = [minmax_score(wide[v], d, CAPS.get(v)) for v, d in items]
        sc[pillar] = pd.concat(parts, axis=1).mean(axis=1, skipna=True)

    # earnings stability: volatility of net margin over the sample (company-level)
    vol = wide.groupby("company").net_margin.std()
    stab = minmax_score(vol, -1)
    sc["Stability"] = sc.company.map(stab)

    if esg_scores is not None and not esg_scores.empty:
        sc = sc.merge(esg_scores, on=["company", "fiscal_year"], how="left")
    pillars = [p for p in list(PILLARS) + ["Stability", "ESG"] if p in sc]
    weights = weights or {p: 1 for p in pillars}
    w = pd.Series({p: weights.get(p, 0) for p in pillars}, dtype=float)
    vals = sc[pillars]
    mask = vals.notna()
    sc["Composite"] = (vals.fillna(0) * w).sum(axis=1) / (mask * w).sum(axis=1).replace(0, np.nan)
    sc["Risk band"] = pd.cut(sc.Composite, [-1, 40, 65, 101], labels=["High risk", "Moderate risk", "Low risk"])
    return sc


# --------------------------------------------------------------------------------------
# ESG data (filled in by the group from integrated / sustainability reports)
# --------------------------------------------------------------------------------------
ESG_COLUMNS = ["company", "fiscal_year", "pillar", "indicator", "value", "unit", "better_when",
               "capital", "source_document", "report_page"]


def load_esg(file) -> pd.DataFrame:
    df = pd.read_csv(file)
    missing = [c for c in ESG_COLUMNS if c not in df.columns]
    if missing:
        raise ValueError(f"ESG file is missing columns: {', '.join(missing)}")
    df["value"] = pd.to_numeric(df["value"], errors="coerce")
    df = df.dropna(subset=["value"])
    df["fiscal_year"] = df.fiscal_year.astype(int)
    return df


def esg_scores(esg: pd.DataFrame) -> pd.DataFrame:
    """ESG pillar score 0-100 per company-year: average of min-max scaled indicators."""
    if esg.empty:
        return pd.DataFrame(columns=["company", "fiscal_year", "ESG"])
    parts = []
    for ind, g in esg.groupby("indicator"):
        direction = -1 if str(g.better_when.iloc[0]).strip().lower().startswith("low") else 1
        g = g.copy()
        g["score"] = minmax_score(g.value, direction)
        parts.append(g)
    allp = pd.concat(parts)
    return allp.groupby(["company", "fiscal_year"]).score.mean().rename("ESG").reset_index()


if __name__ == "__main__":
    lg = build_long()
    wd = build_wide(lg)
    pd.set_option("display.width", 250)
    print(wd[["company", "fiscal_year", "revenue", "net_profit", "total_assets", "total_equity",
              "total_liabilities", "operating_cash_flow", "capex", "total_debt"]].round(0).to_string())
    print(wd[["company", "fiscal_year"] + RATIOS].round(1).to_string())
    print(quality_checks(wd, lg).to_string())
    print(scorecard(wd).round(1).to_string())
